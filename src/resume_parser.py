
from pathlib import Path
from pypdf import PdfReader
from docx import Document


def extract_resume_text(file_path):
    """Extract text from a PDF or DOCX resume."""

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"Resume file not found: {path}"
        )

    extension = path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(str(path))
        text_parts = []

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

        text = "\n".join(text_parts)

    elif extension == ".docx":
        document = Document(str(path))
        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        text = "\n".join(text_parts)

    else:
        raise ValueError(
            "Unsupported file format. Please use PDF or DOCX."
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. The PDF may be scanned "
            "or the document may be empty."
        )

    return text.strip()
